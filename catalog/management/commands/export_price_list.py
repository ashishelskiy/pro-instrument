# catalog/management/commands/export_price_list.py
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from django.core.management.base import BaseCommand
from catalog.models import Product


class Command(BaseCommand):
    help = 'Экспорт товаров в Excel (прайс-лист)'

    def add_arguments(self, parser):
        parser.add_argument(
            '--file',
            type=str,
            default='price_list.xlsx',
            help='Имя файла для экспорта'
        )

    def handle(self, *args, **options):
        filename = options['file']
        products = Product.objects.all().select_related('brand', 'category')

        # Создаем рабочую книгу
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Прайс-лист'

        # Заголовки
        headers = [
            'ID', 'Артикул', 'Название', 'Цена', 'Старая цена',
            'Наличие', 'Бренд', 'Категория', 'Ссылка на фото'
        ]
        ws.append(headers)

        # Стили для заголовков
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='BD0000', end_color='BD0000', fill_type='solid')
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Данные
        for product in products:
            ws.append([
                product.id,
                product.article,
                product.name,
                product.price,
                product.old_price or '',
                'В наличии' if product.in_stock else 'Под заказ',
                product.brand.name if product.brand else '',
                product.category.name if product.category else '',
                product.image if product.image else '',
            ])

        # Автоширина колонок
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

        # Сохраняем
        wb.save(filename)

        self.stdout.write(
            self.style.SUCCESS(f'✅ Экспортировано {products.count()} товаров в {filename}')
        )